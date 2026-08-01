"""Tests for scoresheet parsing in scripts/log_eval_to_wandb.py.

The Fixed and Random tabs do NOT share a column layout: Random carries an extra
"Gripper Pos" column, so every scoring column sits one to the right. These tests
score both tabs through the writer in scripts/make_eval_scoresheet.py, so the
generator and the parser cannot drift apart silently.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

import log_eval_to_wandb as lew  # noqa: E402
import make_eval_scoresheet as mes  # noqa: E402

# trial -> (grasped, placed, failure mode, placement error cm)
TRIALS = [
    ("Y", "Y", None, 2.0),
    ("Y", "N", "grasped-dropped", None),
    ("N", "N", "no-grasp", None),
    ("Y", "Y", None, 4.0),
]


def _write_sheet(path, tab):
    """Write a scoresheet whose `tab` holds TRIALS, using the real generator."""
    cols = mes.TABS[tab]
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet(tab)
    mes.build_tab(ws, "TestPolicy", cols, ["(0,0)"] * mes.N_TRIALS)

    for i, (grasped, placed, failure, err) in enumerate(TRIALS):
        row = mes.FIRST_TRIAL_ROW + i
        ws[f"{cols['grasped']}{row}"] = grasped
        ws[f"{cols['placed']}{row}"] = placed
        if failure is not None:
            ws[f"{cols['failure']}{row}"] = failure
        if err is not None:
            ws[f"{cols['err']}{row}"] = err

    wb.save(path)
    return path


@pytest.mark.parametrize("tab", ["Fixed", "Random"])
def test_compute_metrics_reads_the_tabs_own_columns(tmp_path, tab):
    xlsx = _write_sheet(tmp_path / f"{tab}.xlsx", tab)

    m = lew.compute_metrics(xlsx, tab)

    assert m["trials_scored"] == 4, "unscored rows 5-50 must be skipped"
    assert m["successes"] == 2
    assert m["success_rate"] == pytest.approx(0.5)
    assert m["grasp_rate"] == pytest.approx(0.75)
    assert m["mean_placement_error_cm"] == pytest.approx(3.0)
    assert m["placement_mse_cm2"] == pytest.approx(10.0)
    assert m["fail_no_grasp"] == 1
    assert m["fail_grasped_dropped"] == 1
    assert m["fail_wrong_placement"] == 0
    assert m["fail_other"] == 0


def test_empty_sheet_scores_nothing(tmp_path):
    xlsx = tmp_path / "blank.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    mes.build_tab(wb.create_sheet("Fixed"), "TestPolicy", mes.TABS["Fixed"],
                  ["(0,0)"] * mes.N_TRIALS)
    wb.save(xlsx)

    m = lew.compute_metrics(xlsx, "Fixed")

    assert m["trials_scored"] == 0
    assert m["success_rate"] == 0.0
    assert m["mean_placement_error_cm"] is None


def test_generator_and_parser_agree_on_columns():
    for tab, cols in mes.TABS.items():
        assert lew.TAB_COLUMNS[tab] == {k: cols[k]
                                        for k in ("grasped", "placed", "failure", "err")}
    assert (lew.FIRST_TRIAL_ROW, lew.LAST_TRIAL_ROW) == (mes.FIRST_TRIAL_ROW, mes.LAST_TRIAL_ROW)
